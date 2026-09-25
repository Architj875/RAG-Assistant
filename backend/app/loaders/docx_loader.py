import re

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.loaders.base_loader import BaseLoader
from app.loaders.document_assembler import Block, assemble_documents, table_to_markdown

HEADING_STYLE_PATTERN = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


class DOCXLoader(BaseLoader):
    """
    Word documents carry real structure - paragraph style names
    ("Heading 1", "Heading 2", "Title", "Normal") - unlike a PDF,
    where headings have to be inferred from font size. There's no
    guessing needed here: a paragraph's style name says directly
    whether it's a heading and at what level, so this loader just
    reads that and hands off Blocks, same as every other loader.
    """

    def load(self, file_path: str):
        doc = docx.Document(file_path)

        blocks = []
        position = 0

        for item in self._iter_block_items(doc):
            if isinstance(item, Paragraph):
                text = item.text.strip()
                if not text:
                    continue

                level = self._heading_level(item)
                if level:
                    blocks.append(
                        Block(type="heading", text=text, level=level, page=position)
                    )
                else:
                    blocks.append(Block(type="text", text=text, page=position))

            elif isinstance(item, Table):
                rows = [
                    [cell.text for cell in row.cells] for row in item.rows
                ]
                md = table_to_markdown(rows)
                if md:
                    blocks.append(
                        Block(type="table", text=f"\n{md}\n", page=position)
                    )

            position += 1

        return assemble_documents(blocks, source=file_path)

    @staticmethod
    def _heading_level(paragraph: Paragraph):
        style_name = (paragraph.style.name or "").strip()

        if style_name.lower() == "title":
            return 1

        match = HEADING_STYLE_PATTERN.match(style_name)
        if match:
            # cap at level 3 to match the heading_path depth used
            # elsewhere in the pipeline (PDF loader does the same)
            return min(int(match.group(1)), 3)

        return 0

    @staticmethod
    def _iter_block_items(document: DocxDocument):
        """
        Walk paragraphs and tables in the order they actually appear
        in the document. python-docx doesn't expose this directly -
        doc.paragraphs and doc.tables are separate flat lists with no
        position info - so this reads the underlying XML body in
        document order instead.
        """
        body = document.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                yield Paragraph(child, document)
            elif child.tag == qn("w:tbl"):
                yield Table(child, document)